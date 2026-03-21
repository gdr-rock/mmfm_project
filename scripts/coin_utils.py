#!/usr/bin/env python3
"""Shared utilities for building COIN planning datasets."""

from __future__ import annotations

import csv
import json
import os
import random
import re
from pathlib import Path
from typing import Dict, Iterable, List, Tuple


def humanize_task_name(name: str) -> str:
    """Convert COIN class labels like ReplaceCDDriveWithSSD to readable text."""
    if not name:
        return name
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", name)
    text = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", text)
    return " ".join(text.split())


def build_goal_text(task_name: str) -> str:
    return f"Complete the task: {humanize_task_name(task_name)}."


def read_coin_database(coin_json_path: str) -> Dict[str, dict]:
    """Load COIN video metadata and sort annotation segments by start time."""
    data = json.loads(Path(coin_json_path).read_text())["database"]
    for record in data.values():
        record["annotation"] = sorted(
            record.get("annotation", []),
            key=lambda ann: (float(ann["segment"][0]), float(ann["segment"][1])),
        )
    return data


def _load_taxonomy_xlsx(taxonomy_xlsx_path: str) -> Tuple[Dict[int, str], Dict[int, list]]:
    """Parse COIN taxonomy.xlsx using openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(taxonomy_xlsx_path, read_only=True, data_only=True)
    ws = wb["target_action_mapping"]

    task_names: Dict[int, str] = {}
    task_actions: Dict[int, list] = {}

    first = True
    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        if row is None or len(row) < 4:
            continue
        target_id, target_label, action_id, action_label = row[:4]
        if target_id is None or target_label is None or action_id is None or action_label is None:
            continue

        recipe_type = int(target_id)
        act_id = int(action_id)
        task_label = str(target_label).strip()
        act_label = str(action_label).strip()

        task_names[recipe_type] = task_label
        task_actions.setdefault(recipe_type, []).append(
            {"action_id": act_id, "action": act_label}
        )

    for recipe_type, actions in task_actions.items():
        actions.sort(key=lambda item: item["action_id"])
        for idx, item in enumerate(actions, start=1):
            item["step_idx"] = idx

    return task_names, task_actions


def build_taxonomy_cache(coin_json_path: str, taxonomy_xlsx_path: str) -> dict:
    """Build a machine-readable taxonomy cache for COIN."""
    database = read_coin_database(coin_json_path)
    task_names, task_actions = _load_taxonomy_xlsx(taxonomy_xlsx_path)

    tasks = {}
    step_id_to_action = {}
    for recipe_type, canonical_steps in task_actions.items():
        task_name = task_names[recipe_type]
        tasks[str(recipe_type)] = {
            "task_id": str(recipe_type),
            "recipe_type": recipe_type,
            "task_name": task_name,
            "task_name_human": humanize_task_name(task_name),
            "goal": build_goal_text(task_name),
            "canonical_steps": canonical_steps,
        }
        for step in canonical_steps:
            step_id_to_action[str(step["action_id"])] = step["action"]

    missing = []
    for record in database.values():
        recipe_type = int(record["recipe_type"])
        if str(recipe_type) not in tasks:
            missing.append(recipe_type)
            continue
        task = tasks[str(recipe_type)]
        if record["class"] != task["task_name"]:
            raise ValueError(
                f"Task name mismatch for recipe_type={recipe_type}: "
                f"COIN.json='{record['class']}' vs taxonomy='{task['task_name']}'"
            )
        known_ids = {step["action_id"] for step in task["canonical_steps"]}
        for ann in record.get("annotation", []):
            ann_id = int(ann["id"])
            if ann_id not in known_ids:
                raise ValueError(
                    f"Missing step id {ann_id} for recipe_type={recipe_type} in taxonomy"
                )

    if missing:
        raise ValueError(f"Missing taxonomy entries for recipe_type values: {sorted(set(missing))[:20]}")

    return {
        "tasks": tasks,
        "step_id_to_action": step_id_to_action,
    }


def ensure_taxonomy_cache(
    coin_json_path: str,
    taxonomy_xlsx_path: str,
    cache_path: str,
    force: bool = False,
) -> dict:
    """Load taxonomy cache if present; otherwise build and persist it."""
    cache_file = Path(cache_path)
    if cache_file.exists() and not force:
        return json.loads(cache_file.read_text())

    cache = build_taxonomy_cache(coin_json_path, taxonomy_xlsx_path)
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(json.dumps(cache, indent=2, ensure_ascii=False) + "\n")
    return cache


def load_taxonomy_cache(cache_path: str) -> dict:
    return json.loads(Path(cache_path).read_text())


def build_video_split_rows(
    database: Dict[str, dict],
    taxonomy: dict,
    val_frac: float,
    seed: int,
) -> List[dict]:
    """Create a deterministic train/val/test video split manifest."""
    rng = random.Random(seed)
    rows: List[dict] = []

    by_task_training: Dict[str, List[Tuple[str, dict]]] = {}
    for video_id, record in database.items():
        task_id = str(int(record["recipe_type"]))
        split = "test" if record["subset"] == "testing" else None
        if split:
            rows.append(_build_split_row(video_id, record, taxonomy, split))
        else:
            by_task_training.setdefault(task_id, []).append((video_id, record))

    for task_id, videos in by_task_training.items():
        rng.shuffle(videos)
        n_val = max(1, int(len(videos) * val_frac))
        val_ids = {video_id for video_id, _ in videos[:n_val]}
        for video_id, record in videos:
            split = "val" if video_id in val_ids else "train"
            rows.append(_build_split_row(video_id, record, taxonomy, split))

    rows.sort(key=lambda row: (int(row["task_id"]), row["split"], row["video_id"]))
    return rows


def _build_split_row(video_id: str, record: dict, taxonomy: dict, split: str) -> dict:
    task_id = str(int(record["recipe_type"]))
    task = taxonomy["tasks"][task_id]
    return {
        "task_id": task_id,
        "task_name": task["task_name"],
        "video_id": video_id,
        "split": split,
        "official_subset": record["subset"],
        "recipe_type": int(record["recipe_type"]),
        "video_url": record["video_url"],
        "roi_start": float(record["start"]),
        "roi_end": float(record["end"]),
        "duration": float(record["duration"]),
    }


def load_split_manifest(split_csv_path: str) -> Dict[str, dict]:
    """Load split rows keyed by video_id."""
    mapping = {}
    with open(split_csv_path) as handle:
        for row in csv.DictReader(handle):
            mapping[row["video_id"]] = row
    return mapping


def format_system1_prompt(
    goal: str,
    interpretation: str,
    prefix_steps: List[Tuple[str, str]],
    k: int,
) -> str:
    """Build a System-1 prompt that matches standalone inference formatting."""
    lines = [f"Goal: {goal}"]
    if interpretation:
        lines.append(f"Interpretation: {interpretation}")
    lines.append("")
    lines.append("Progress so far:")
    if prefix_steps:
        for idx, (action, state_change) in enumerate(prefix_steps, start=1):
            lines.append(f"  {idx}) {action} | {state_change}")
    else:
        lines.append("  (No steps observed yet.)")
    lines.append("")
    lines.append(
        f'Predict the next {k} step(s). Output JSON only: '
        '{"next_steps": [{"action": "...", "state_change": "..."}, ...]}'
    )
    return "\n".join(lines)


def format_system1_output(target_steps: List[Tuple[str, str]]) -> str:
    payload = {
        "next_steps": [
            {"action": action, "state_change": state_change}
            for action, state_change in target_steps
        ]
    }
    return json.dumps(payload, ensure_ascii=False)


def make_interpretation(task_steps: List[str]) -> str:
    """Generate the rule-based interpretation sentence."""
    if task_steps:
        return (
            f"The task begins with '{task_steps[0]}' and is considered complete "
            f"once '{task_steps[-1]}' is done."
        )
    return "The task is carried out from the current state to the final goal state."


def trajectory_from_transition_rows(rows: List[dict]) -> List[Tuple[str, str, int]]:
    """Reconstruct raw observed steps from transition rows."""
    ordered = sorted(rows, key=lambda row: int(row["seg_pos"]))
    steps = [
        (row["action"], row["state_change"], int(row["canonical_order"]))
        for row in ordered
    ]
    last = ordered[-1]
    steps.append(
        (
            last["next_action"],
            last["next_state_change"],
            int(last["next_canonical_order"]),
        )
    )
    return steps


def canonical_remaining_steps(task: dict, frontier: int) -> List[Tuple[str, str]]:
    """Return canonical steps after the frontier with generated state changes."""
    remaining = []
    for step in task["canonical_steps"]:
        if int(step["step_idx"]) <= frontier:
            continue
        action = step["action"]
        remaining.append((action, step_to_state_change(action)))
    return remaining


def prefix_frontier(prefix_steps: Iterable[Tuple[str, str, int]]) -> int:
    frontier = 0
    for _, _, canonical_order in prefix_steps:
        frontier = max(frontier, int(canonical_order))
    return frontier


def load_transition_rows(csv_path: str) -> Tuple[Dict[Tuple[str, str], List[dict]], Dict[str, dict]]:
    """Load transition rows grouped by (task_id, video_id)."""
    grouped: Dict[Tuple[str, str], List[dict]] = {}
    task_info: Dict[str, dict] = {}
    with open(csv_path) as handle:
        for row in csv.DictReader(handle):
            key = (row["task_id"], row["video_id"])
            grouped.setdefault(key, []).append(row)
            task_info.setdefault(
                row["task_id"],
                {
                    "name": row["task_name"],
                    "goal": row["task_goal"],
                },
            )
    return grouped, task_info


def _cap(text: str) -> str:
    return text[0].upper() + text[1:] if text else text


_EXACT_OVERRIDES = {
    "pull up the hair to reserve place for the hair extensions":
        "The hair is now lifted to make space for the extensions.",
    "put on the hair extensions":
        "The hair extensions are now attached.",
    "put down the hair and comb":
        "The hair is now let down and combed into place.",
    "tighten screws":
        "The screws are now tightened securely.",
    "jack up the car":
        "The vehicle is now jacked up off the ground.",
    "remove the tire":
        "The tire is now removed from the vehicle.",
    "put on the tire":
        "The replacement tire is now mounted on the vehicle.",
    "unscrew the screw":
        "The screw is now loosened and removed.",
}

_VERB_TEMPLATES = {
    "add": lambda obj: f"{_cap(obj)} is now added to the mixture.",
    "apply": lambda obj: f"{_cap(obj)} is now applied.",
    "assemble": lambda obj: f"{_cap(obj)} is now assembled.",
    "attach": lambda obj: f"{_cap(obj)} is now attached and secured.",
    "begin": lambda obj: f"{_cap(obj)} is now started.",
    "blend": lambda obj: f"{_cap(obj)} is now blended together.",
    "boil": lambda obj: f"{_cap(obj)} is now boiled.",
    "brush": lambda obj: f"{_cap(obj)} is now brushed.",
    "clean": lambda obj: f"{_cap(obj)} is now cleaned.",
    "close": lambda obj: f"{_cap(obj)} is now closed.",
    "comb": lambda obj: f"{_cap(obj)} is now combed.",
    "connect": lambda obj: f"{_cap(obj)} is now connected.",
    "cook": lambda obj: f"{_cap(obj)} is now cooked.",
    "cut": lambda obj: f"{_cap(obj)} is now cut into pieces.",
    "drain": lambda obj: f"{_cap(obj)} is now drained.",
    "dry": lambda obj: f"{_cap(obj)} is now dried.",
    "fasten": lambda obj: f"{_cap(obj)} is now fastened securely.",
    "fill": lambda obj: f"{_cap(obj)} is now filled.",
    "fix": lambda obj: f"{_cap(obj)} is now fixed in place.",
    "fold": lambda obj: f"{_cap(obj)} is now folded.",
    "heat": lambda obj: f"{_cap(obj)} is now heated.",
    "insert": lambda obj: f"{_cap(obj)} is now inserted into position.",
    "install": lambda obj: f"{_cap(obj)} is now installed.",
    "jack": lambda obj: f"{_cap(obj) if obj else 'The vehicle'} is now raised.",
    "make": lambda obj: f"{_cap(obj)} is now made.",
    "mix": lambda obj: f"{_cap(obj) if obj else 'The ingredients'} are now mixed together.",
    "mount": lambda obj: f"{_cap(obj)} is now mounted.",
    "open": lambda obj: f"{_cap(obj)} is now opened.",
    "peel": lambda obj: f"{_cap(obj)} is now peeled.",
    "place": lambda obj: f"{_cap(obj)} is now placed in position.",
    "pour": lambda obj: f"{_cap(obj)} is now poured in.",
    "press": lambda obj: f"{_cap(obj)} is now pressed down.",
    "pull": lambda obj: f"{_cap(obj)} is now pulled into place.",
    "put": lambda obj: f"{_cap(obj)} is now placed in position.",
    "remove": lambda obj: f"{_cap(obj)} is now removed.",
    "replace": lambda obj: f"{_cap(obj)} is now replaced.",
    "run": lambda obj: f"{_cap(obj)} is now in motion.",
    "screw": lambda obj: f"{_cap(obj)} is now screwed on tight.",
    "set": lambda obj: f"{_cap(obj)} is now set into place.",
    "shake": lambda obj: f"{_cap(obj)} is now shaken.",
    "slice": lambda obj: f"{_cap(obj)} is now sliced.",
    "spread": lambda obj: f"{_cap(obj)} is now spread evenly.",
    "stir": lambda obj: f"{_cap(obj) if obj else 'The contents'} are now stirred and combined.",
    "take": lambda obj: f"{_cap(obj)} is now taken out.",
    "tighten": lambda obj: f"{_cap(obj)} is now tightened securely.",
    "turn": lambda obj: f"{_cap(obj)} is now turned into place.",
    "wash": lambda obj: f"{_cap(obj)} is now washed.",
    "whisk": lambda obj: f"{_cap(obj) if obj else 'The mixture'} is now whisked until smooth.",
    "wrap": lambda obj: f"{_cap(obj)} is now wrapped.",
}


def step_to_state_change(step_text: str) -> str:
    """Convert a procedural step string into a coarse state-change sentence."""
    step = step_text.strip().lower()
    if step in _EXACT_OVERRIDES:
        return _EXACT_OVERRIDES[step]

    for verb in sorted(_VERB_TEMPLATES.keys(), key=len, reverse=True):
        if step == verb:
            return _VERB_TEMPLATES[verb]("")
        if step.startswith(verb + " "):
            obj = step[len(verb):].strip()
            return _VERB_TEMPLATES[verb](obj)

    return f"{_cap(step)} is now done." if step else "The step is now done."
